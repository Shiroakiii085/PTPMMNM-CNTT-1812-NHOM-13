# -*- coding: utf-8 -*-
"""
Script tạo tài liệu Huong_Dan_Demo_Bao_Cao.docx chuyên nghiệp
dành cho báo cáo đồ án Đề tài 4 trước giảng viên.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, title="LƯU Ý KHI THUYẾT TRÌNH:"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F9FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="0284C7"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"📌 {title} ")
    r1.bold = True
    r1.font.color.rgb = RGBColor(3, 105, 161)
    r1.font.size = Pt(10.5)
    
    r2 = p.add_run(text)
    r2.font.color.rgb = RGBColor(15, 23, 42)
    r2.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_if_exists(doc, img_path, caption, width=5.5):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width))
        
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_after = Pt(8)
        cap_run = cap_p.add_run(f"Hình: {caption}")
        cap_run.font.italic = True
        cap_run.font.size = Pt(9.5)
        cap_run.font.color.rgb = RGBColor(100, 116, 139)

def build_demo_document():
    doc = Document()
    
    # Cấu hình lề trang (1 inch = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    img_dir = r"C:\Users\hpham\.gemini\antigravity-ide\brain\00005028-e854-4434-8c4e-d0a25e279a0e"
    
    # -------------------------------------------------------------------------
    # TRANG BÌA & TIÊU ĐỀ
    # -------------------------------------------------------------------------
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_header.add_run("TRƯỜNG ĐẠI HỌC ĐẠI NAM - KHOA CÔNG NGHỆ THÔNG TIN\n")
    r_uni.bold = True
    r_uni.font.size = Pt(12)
    r_uni.font.color.rgb = RGBColor(71, 85, 105)
    
    r_sub = p_header.add_run("HỌC PHẦN: PHÁT TRIỂN PHẦN MỀM MÃ NGUỒN MỞ\n")
    r_sub.bold = True
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_main_title = p_title.add_run("TÀI LIỆU HƯỚNG DẪN KỊCH BẢN DEMO & BÁO CÁO BẢO VỆ ĐỒ ÁN\n")
    r_main_title.bold = True
    r_main_title.font.size = Pt(16)
    r_main_title.font.color.rgb = RGBColor(15, 23, 42)
    
    r_topic = p_title.add_run("ĐỀ TÀI 4: XÂY DỰNG CÔNG CỤ KIỂM THỬ TỰ ĐỘNG QUYỀN TRUY CẬP TRONG ỨNG DỤNG WEB (ACCESSTESTER)\n")
    r_topic.bold = True
    r_topic.font.size = Pt(14)
    r_topic.font.color.rgb = RGBColor(2, 132, 199)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Sinh viên thực hiện: Nhóm 13 | Hệ Thống Bệnh Án Điện Tử MediRecord & AccessTester v2.0")
    r_meta.font.italic = True
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------------------
    # PHẦN I: TỔNG QUAN HỆ THỐNG & NHỮNG GÌ ĐÃ HOÀN THÀNH
    # -------------------------------------------------------------------------
    h1 = doc.add_heading("I. TỔNG QUAN HỆ THỐNG & NHỮNG THÀNH TỰU ĐÃ ĐẠT ĐƯỢC", level=1)
    h1.runs[0].font.color.rgb = RGBColor(3, 105, 161)
    
    p_intro = doc.add_paragraph(
        "Đề tài giải quyết bài toán an toàn thông tin cấp thiết: 'Một trang Web có thể khai báo phân quyền đúng "
        "nhưng trong mã nguồn vẫn để lộ dữ liệu nhạy cảm của người này cho người khác xem (Lỗ hổng BOLA / IDOR - CWE-639)'. "
        "Nhóm 13 đã xây dựng hoàn chỉnh hệ sinh thái kiểm thử độc lập gồm 2 thành phần phần mềm chuyên nghiệp:"
    )
    p_intro.paragraph_format.space_after = Pt(6)
    
    # Bảng 2 thành phần
    tbl_comp = doc.add_table(rows=3, cols=3)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Thành phần", "Địa chỉ / Cổng", "Mô tả vai trò & Tính năng cốt lõi"]
    for i, h in enumerate(headers):
        cell = tbl_comp.cell(0, i)
        set_cell_background(cell, "0284C7")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)
        
    data_comp = [
        ("Ứng dụng Mục tiêu (Target App)\nMediRecord EMR", "http://localhost:8205\n(127.0.0.1:8205)", 
         "- Hệ thống quản lý bệnh án điện tử y tế thực tế.\n"
         "- Bác sĩ Lê Hoàng Minh gửi bệnh án cho Bệnh nhân An, Bệnh nhân Bình.\n"
         "- Thanh chuyển đổi tài khoản 1-Click Fast Switch Bar.\n"
         "- Hộp Tra Cứu & Can Thiệp Bệnh Án Theo ID Tự Do (Test IDOR trực tiếp cả Xem và Sửa).\n"
         "- Nút tròn nổi FAB Terminal ở góc dưới bên phải màn hình điều khiển trạng thái lỗi.\n"
         "- Module Quản Trị Người Dùng Toàn Diện (Thêm, Sửa, Xóa, Tìm kiếm) cho Quản trị viên.\n"
         "- Trang trả về mã lỗi bảo mật HTTP 403 Forbidden chuyên biệt."),
        ("Công cụ Kiểm thử Tự động\nAccessTester Engine & Web UI", "http://localhost:8105\n(127.0.0.1:8105)", 
         "- Quản trị ma trận phân quyền RBAC & ABAC (Ownership scope).\n"
         "- Khai báo và thực thi tự động 24 ca kiểm thử HTTP chuẩn.\n"
         "- Cơ chế phân loại 3 trạng thái: PASS, FAIL (Lỗ hổng), ERROR (Mạng/kỹ thuật).\n"
         "- Bộ lọc địa chỉ an toàn (Target Allowlist) chống tấn công SSRF.\n"
         "- Cơ chế Redactor tự động khử thông tin nhạy cảm (token, mật khẩu).\n"
         "- Giao diện Đối chiếu 2 Đợt Chạy (Compare Runs) kèm biểu đồ Chart.js trực quan.")
    ]
    
    for row_idx, data in enumerate(data_comp, start=1):
        for col_idx, text in enumerate(data):
            cell = tbl_comp.cell(row_idx, col_idx)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------------------
    # PHẦN II: CHUẨN BỊ MÔI TRƯỜNG TRƯỚC KHI BẮT ĐẦU DEMO
    # -------------------------------------------------------------------------
    h2 = doc.add_heading("II. DANH MỤC CHUẨN BỊ MÔI TRƯỜNG TRƯỚC BUỔI BÁO CÁO", level=1)
    h2.runs[0].font.color.rgb = RGBColor(3, 105, 161)

    doc.add_paragraph(
        "Để buổi báo cáo diễn ra trơn tru nhất, trước khi giảng viên yêu cầu demo, sinh viên thực hiện mở sẵn "
        "các cửa sổ dòng lệnh và khởi động 2 dịch vụ nền:"
    )
    
    add_callout(
        doc,
        "Nên mở sẵn 2 tab trình duyệt: Tab 1 tại http://localhost:8205 (MediRecord EMR) và Tab 2 tại "
        "http://localhost:8105 (AccessTester). Điều này giúp thao tác trình diễn diễn ra liền mạch, không bị gián đoạn.",
        title="MẸO PHÒNG BÁO CÁO:"
    )

    steps_prep = [
        ("Bước 1", "Khởi động CSDL PostgreSQL (WSL/Linux):", "sudo service postgresql start"),
        ("Bước 2", "Nạp cấu hình và dữ liệu mẫu sạch sẽ (1 lệnh duy nhất):", "cd /mnt/c/University/Nam_3_K1/PTPMMNM/btvn/lab5/Ma_nguon\nbash lab.sh setup"),
        ("Bước 3", "Kiểm tra tình trạng môi trường:", "bash lab.sh check\n(Kết quả in ra: MOI_TRUONG=SAN_SANG)"),
        ("Bước 4", "Khởi động Ứng dụng Y tế MediRecord (Cổng 8205):", "bash lab.sh target:start\n(Hoặc: php -S 127.0.0.1:8205 -t sample-app/public)"),
        ("Bước 5", "Khởi động Công cụ AccessTester Web UI (Cổng 8105):", "php -S 127.0.0.1:8105 -t public")
    ]
    
    for s_step, s_title, s_cmd in steps_prep:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_step = p.add_run(f"• {s_step}: {s_title}\n")
        r_step.bold = True
        r_step.font.color.rgb = RGBColor(15, 23, 42)
        r_cmd = p.add_run(f"   {s_cmd}")
        r_cmd.font.name = "Consolas"
        r_cmd.font.size = Pt(9.5)
        r_cmd.font.color.rgb = RGBColor(30, 64, 175)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------------------
    # PHẦN III: KỊCH BẢN DEMO CHI TIẾT TỪNG BƯỚC (TRÌNH DIỄN TRỌN VẸN)
    # -------------------------------------------------------------------------
    h3 = doc.add_heading("III. KỊCH BẢN TRÌNH DIỄN DEMO TỪNG BƯỚC (STEP-BY-STEP)", level=1)
    h3.runs[0].font.color.rgb = RGBColor(3, 105, 161)

    doc.add_paragraph(
        "Kịch bản thuyết trình được thiết kế tối ưu trong thời lượng 8 - 10 phút, "
        "làm nổi bật toàn bộ tính năng kỹ thuật và trực quan hoá cao nhất cho giảng viên:"
    )

    # Giai đoạn 1
    h3_1 = doc.add_heading("Giai đoạn 1 (0:00 - 1:30): Mở đầu, Đặt vấn đề & Giới thiệu Kiến trúc", level=2)
    h3_1.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    p_script1 = doc.add_paragraph()
    r = p_script1.add_run("Lời thoại gợi ý cho sinh viên: ")
    r.bold = True
    r.font.italic = True
    p_script1.add_run(
        "\"Kính thưa Thầy/Cô, trong các hệ thống Web hiện đại, phân quyền vai trò (RBAC) thường được khai báo chặt chẽ, "
        "tuy nhiên lỗi phân quyền cấp đối tượng (BOLA / IDOR) vẫn là lỗ hổng an ninh mạng đứng số 1 theo OWASP API Security. "
        "Một người dùng đã đăng nhập hợp lệ nhưng nếu hệ thống không kiểm tra mối quan hệ sở hữu bản ghi thì người đó hoàn toàn "
        "có thể xem hoặc đánh cắp dữ liệu riêng tư của người khác. Nhóm 13 đã giải quyết trọn vẹn bài toán này bằng việc xây dựng: "
        "Thứ nhất là ứng dụng y tế MediRecord (cổng 8205) để mô phỏng dữ liệu bệnh án thật; và Thứ hai là công cụ kiểm thử tự động "
        "AccessTester (cổng 8105) có khả năng tự động rà quét, phát hiện rò rỉ và đối chiếu kết quả trước/sau khi vá lỗi.\""
    )

    # Giai đoạn 2
    h3_2 = doc.add_heading("Giai đoạn 2 (1:30 - 5:00): Trình diễn Ứng dụng Y tế MediRecord EMR (Cổng 8205)", level=2)
    h3_2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Mở tab trình duyệt tại http://localhost:8205 để trình diễn các nghiệp vụ và trực quan hoá lỗi bảo mật:"
    )

    demo_steps_8205 = [
        ("Bước 2.1: Giới thiệu giao diện, Fast Switch Bar và Nút tròn nổi FAB Terminal",
         "Chỉ vào thanh đen Fast Switch Bar trên cùng: Giới thiệu hệ thống có 5 tài khoản thật (Bác sĩ Minh, Bệnh nhân An, "
         "Bệnh nhân Bình, Quản trị viên và Kỹ sư Kiểm thử). Nhấn thử vào từng nút để chuyển đổi tức thì góc nhìn người dùng.\n"
         "Đặc biệt, chỉ vào Nút tròn nổi FAB ở góc dưới bên phải màn hình (Floating Action Button với biểu tượng '>_'). "
         "Bấm vào nút tròn: Một cửa sổ Terminal trượt lên ngay tại chỗ giúp điều khiển trạng thái cờ lỗi và chạy lệnh trực tiếp."),
        
        ("Bước 2.2: Bác sĩ Lê Hoàng Minh lập & gửi bệnh án mới cho Bệnh nhân",
         "Bấm chọn 'Bác sĩ Minh (Doctor)' -> Nhấn nút 'Bác Sĩ Gửi Bệnh Án' trên menu. Chọn người nhận là Bệnh nhân An, "
         "nhập chẩn đoán và bấm 'Lập & Gửi Bệnh Án'. Hồ sơ mới được ghi nhận ngay lập tức vào cơ sở dữ liệu."),
        
        ("Bước 2.3: Bật lỗi IDOR bằng Nút tròn nổi FAB Terminal",
         "Bấm vào Nút tròn nổi FAB ở góc dưới bên phải màn hình -> Cửa sổ Drawer Terminal mở ra -> Bấm nút 'Lệnh: Tạo Lỗi (set BUG_IDOR=1)'. "
         "Hệ thống cập nhật ngay lập tức: Đèn tín hiệu trên nút tròn chuyển sang màu đỏ và thanh thông báo báo đỏ rực: 'LỖI IDOR: ĐANG BẬT'."),

        ("Bước 2.4: Trực quan hóa Lỗ hổng IDOR qua Hộp Tra Cứu & Can Thiệp Theo Mã ID Tự Do",
         "1. Chuyển sang góc nhìn 'Bệnh nhân An (Patient)' (An chỉ sở hữu hồ sơ số 01 và 02).\n"
         "2. Tại Hộp Tra Cứu 'TRA CỨU & MỞ HỒ SƠ BỆNH ÁN THEO MÃ ID', sinh viên nhập tự do số '3' (Mã bệnh án ung bướu của Bình) và bấm 'Tra Cứu Bệnh Án'.\n"
         "3. KẾT QUẢ XEM TRÁI PHÉP: Màn hình hiển thị toàn bộ hồ sơ bệnh án ung bướu tuyệt mật của Bình kèm Hộp Cảnh Báo Bảo Mật Đỏ Rực "
         "thông báo: Bệnh nhân An đang xem trái phép dữ liệu của Bệnh nhân Bình!\n"
         "4. KẾT QUẢ SỬA TRÁI PHÉP (IDOR on Update): An tiếp tục bấm nút 'Chỉnh Sửa Bệnh Án' (hoặc truy cập ?view=edit_record&id=3) -> "
         "Mở form sửa đổi chẩn đoán của Bình -> Sửa thành công và hệ thống cảnh báo đỏ rực: '[CẢNH BÁO VI PHẠM BẢO MẬT] KHAI THÁC LỖ HỔNG IDOR / BOLA THÀNH CÔNG!'."),

        ("Bước 2.5: Trực quan hóa Khắc phục Lỗ hổng & Trang Lỗi HTTP 403 Forbidden",
         "1. Bấm vào Nút tròn nổi FAB góc dưới bên phải, nhấn nút 'Lệnh: Vá Lỗi (set BUG_IDOR=0)'. Đèn trạng thái chuyển xanh an toàn.\n"
         "2. Vẫn trong vai Bệnh nhân An, tiếp tục gõ ID '3' và bấm 'Tra Cứu' hoặc 'Chỉnh Sửa'.\n"
         "3. KẾT QUẢ SHOW GIẢNG VIÊN: Hệ thống lập tức chặn đứng và chuyển sang Trang Trả Về Lỗi HTTP 403 Forbidden chuyên biệt, "
         "giải thích cơ chế ABAC từ chối quyền truy cập do không phải chủ sở hữu."),

        ("Bước 2.6: Phân quyền Quản Trị Viên (admin_bv) & Module Quản Trị Người Dùng Toàn Diện",
         "1. Chuyển sang tài khoản 'Quản trị viên (Admin)' -> Bấm vào menu 'Quản Trị Người Dùng' (/?view=users).\n"
         "2. Tìm kiếm tức thì: Gõ từ khóa vào ô tìm kiếm (ví dụ 'Minh' hoặc 'Bác sĩ') -> Danh sách lọc kết quả ngay lập tức mà không cần reload trang.\n"
         "3. Thêm người dùng: Bấm 'Thêm Người Dùng Mới', điền tài khoản, vai trò, khoa phòng -> Lưu thành công vào hệ thống.\n"
         "4. Sửa & Xóa người dùng: Thử bấm nút Sửa thông tin và nút Xóa tài khoản (có hộp thoại xác nhận an toàn, có cơ chế cấm tự xóa chính mình).\n"
         "5. Kiểm soát truy cập: Nếu chuyển sang Bệnh nhân An mà cố tình gõ URL ?view=users, hệ thống sẽ trả về ngay trang lỗi HTTP 403 Forbidden!")
    ]

    for s_title, s_desc in demo_steps_8205:
        p = doc.add_paragraph()
        r_t = p.add_run(f"👉 {s_title}:\n")
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(2, 132, 199)
        p.add_run(s_desc)
        p.paragraph_format.space_after = Pt(4)

    # Chèn ảnh minh họa cổng 8205
    add_image_if_exists(doc, os.path.join(img_dir, "fab_and_lookup_main.png"), "Giao diện chính MediRecord EMR với Hộp tra cứu ID bệnh án tự do và Nút tròn nổi FAB Terminal góc dưới bên phải")
    add_image_if_exists(doc, os.path.join(img_dir, "terminal_fab_drawer.png"), "Cửa sổ Terminal Console trượt lên từ Nút tròn nổi FAB ở góc dưới bên phải màn hình")
    add_image_if_exists(doc, os.path.join(img_dir, "idor_lookup_patient_an.png"), "Bệnh nhân An nhập ID 3 tra cứu và khai thác IDOR thành công kèm cảnh báo đỏ rực rò rỉ dữ liệu y tế")
    add_image_if_exists(doc, os.path.join(img_dir, "idor_edit_record_form.png"), "Giao diện can thiệp chỉnh sửa bệnh án trái phép (IDOR on Update) khi bật lỗi bảo mật")
    add_image_if_exists(doc, os.path.join(img_dir, "patient_an_403_1792814420387.png"), "Trang lỗi bảo mật HTTP 403 Forbidden chuyên biệt khi đã vá lỗi bảo mật")
    add_image_if_exists(doc, os.path.join(img_dir, "doctor_minh_create_form_1792814504352.png"), "Giao diện Bác sĩ Lê Hoàng Minh lập và gửi hồ sơ bệnh án cho bệnh nhân")
    add_image_if_exists(doc, os.path.join(img_dir, "admin_user_crud.png"), "Giao diện Quản Trị Người Dùng toàn diện (Thêm, Sửa, Xóa, Tìm kiếm) dành riêng cho Quản trị viên")
    add_image_if_exists(doc, os.path.join(img_dir, "admin_user_search.png"), "Tính năng tìm kiếm tài khoản người dùng thời gian thực trên giao diện Quản trị viên")

    # Giai đoạn 3
    h3_3 = doc.add_heading("Giai đoạn 3 (5:00 - 8:30): Trình diễn Công cụ Kiểm thử Tự động AccessTester (Cổng 8105)", level=2)
    h3_3.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Chuyển sang tab trình duyệt tại http://localhost:8105 để trình diễn sức mạnh kiểm thử tự động của AccessTester:"
    )

    demo_steps_8105 = [
        ("Bước 3.1: Giới thiệu Quản lý Dự án & Danh mục 24 Ca kiểm thử HTTP",
         "Mở mục 'Dự án' và 'Bộ ca thử': Giải thích ma trận 24 ca kiểm thử bao quát đầy đủ 6 API x 4 ca thử nghiệm (P01-P04 hồ sơ, N01-N04 danh sách bệnh án, "
         "R01-R04 chi tiết bệnh án, U01-U04 cập nhật bệnh án, A01-A04 thông báo chung, M01-M04 quản trị người dùng). "
         "Nhấn mạnh ca kiểm thử R02 chính là ca thử nghiệm bẫy phát hiện IDOR."),

        ("Bước 3.2: Chạy kiểm thử Đợt 1 (Khi ứng dụng còn Lỗi BUG_IDOR=1)",
         "Nhấn nút 'Chạy kiểm thử ngay' (hoặc chọn Dự án #1). Hệ thống gửi 24 yêu cầu HTTP song song.\n"
         "KẾT QUẢ SHOW GIẢNG VIÊN: Đợt chạy ghi nhận: 23 PASS, 1 FAIL, 0 ERROR. "
         "Bấm vào chi tiết ca R02: Hiển thị cờ 'leak_detected = true' do máy chủ trả về HTTP 200 kèm chuỗi nhạy cảm của Bob/Bình."),

        ("Bước 3.3: Chạy kiểm thử Đợt 2 (Sau khi đã vá lỗi BUG_IDOR=0)",
         "Sau khi ứng dụng mục tiêu được vá lỗi, bấm chạy đợt kiểm thử mới.\n"
         "KẾT QUẢ SHOW GIẢNG VIÊN: Đợt chạy ghi nhận hoàn hảo: 24 PASS, 0 FAIL, 0 ERROR (100% ĐẠT CHUẨN). "
         "Ca R02 đã chuyển sang trạng thái PASS vì máy chủ đã chặn đúng mã HTTP 403 Forbidden."),

        ("Bước 3.4: Trình diễn Tính năng 'Đối chiếu 2 Đợt Chạy' (Compare Runs)",
         "Mở menu 'Đối chiếu' tại thanh điều hướng:\n"
         "- Chọn Đợt chạy A: Đợt có lỗi (23 PASS, 1 FAIL).\n"
         "- Chọn Đợt chạy B: Đợt an toàn (24 PASS, 0 FAIL).\n"
         "- Bấm nút 'Đối chiếu 2 đợt chạy'.\n"
         "KẾT QUẢ SHOW GIẢNG VIÊN: Màn hình hiển thị Biểu đồ cột so sánh tỉ lệ PASS/FAIL và bảng chi tiết làm nổi bật duy nhất ca R02 với nhãn: "
         "'[LỖ HỔNG -> ĐÃ KHẮC PHỤC] SỬA THÀNH CÔNG'."),

        ("Bước 3.5: Trình diễn Báo cáo Kiểm thử & Cơ chế Redactor Khử Dữ liệu Nhạy Cảm",
         "Mở Báo cáo kiểm thử: Chỉ cho giảng viên thấy toàn bộ mã Bearer Token và mật khẩu CSDL đã được bộ lọc Redactor tự động băm che chắn "
         "(chỉ hiện dấu sao hoặc mã giả lập), đảm bảo an toàn tuyệt đối khi gửi báo cáo cho các bên liên quan.")
    ]

    for s_title, s_desc in demo_steps_8105:
        p = doc.add_paragraph()
        r_t = p.add_run(f"👉 {s_title}:\n")
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(2, 132, 199)
        p.add_run(s_desc)
        p.paragraph_format.space_after = Pt(4)

    # Chèn ảnh minh họa cổng 8105
    add_image_if_exists(doc, os.path.join(img_dir, "cases_page_1792765519850.png"), "Danh mục 24 ca kiểm thử truy cập được quản lý trực quan trên AccessTester")
    add_image_if_exists(doc, os.path.join(img_dir, "runs_page_1792765622496.png"), "Quản lý lịch sử các đợt chạy kiểm thử tự động")
    add_image_if_exists(doc, os.path.join(img_dir, "comparison_page_1792806205155.png"), "Tính năng Đối chiếu 2 đợt chạy (Compare Runs) chứng minh việc sửa lỗi IDOR thành công")

    # Giai đoạn 4
    h3_4 = doc.add_heading("Giai đoạn 4 (8:30 - 10:00): Trình diễn Kiểm thử Dòng lệnh & Bộ Test Suite 100% PASS", level=2)
    h3_4.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph(
        "Mở terminal WSL / Linux và chạy lệnh kiểm thử toàn diện:\n"
        "bash lab.sh test\n"
        "Chỉ cho giảng viên thấy kết quả chạy tự động qua 3 cấp độ kiểm thử chuyên sâu:"
    )

    p_test_suite = doc.add_paragraph()
    p_test_suite.add_run(
        "1. KIỂM THỬ ĐƠN VỊ (UNIT TESTS): 39/39 bài test PASS (Policy engine, Allowlist SSRF, Classifier, Redactor).\n"
        "2. KIỂM THỬ TÍCH HỢP CSDL (INTEGRATION TESTS): 29/29 bài test PASS (CRUD Dự án, phân quyền Owner/Editor/Viewer, Snapshot, Compare).\n"
        "3. KIỂM THỬ ĐẦU-CUỐI E2E (END-TO-END TESTS): 20/20 bài test PASS (Chứng minh 5 phép thử bắt buộc theo đề cương).\n"
        "=> TỔNG CỘNG: 88 TESTS, 0 FAILED (100% ĐẠT CHUẨN)."
    )
    p_test_suite.paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------------------
    # PHẦN IV: CÂU HỎI PHẢN BIỆN THƯỜNG GẶP CỦA GIẢNG VIÊN & CÁCH TRẢ LỜI
    # -------------------------------------------------------------------------
    h4 = doc.add_heading("IV. BỘ CÂU HỎI PHẢN BIỆN CỦA GIẢNG VIÊN & GỢI Ý CÂU TRẢ LỜI XUẤT SẮC", level=1)
    h4.runs[0].font.color.rgb = RGBColor(3, 105, 161)

    qa_list = [
        ("Câu 1: Lỗ hổng IDOR trong đề tài là gì và nhóm đã sửa trong mã nguồn như thế nào?",
         "Trả lời: IDOR (Insecure Direct Object Reference) hay BOLA xảy ra khi máy chủ chỉ kiểm tra người dùng đã đăng nhập (Authentication) "
         "nhưng bỏ quên bước kiểm tra quan hệ sở hữu đối tượng (Authorization). Trong mã nguồn 'sample-app/public/index.php', ở bản lỗi, "
         "khi nhận được id=3, server trả về ngay dữ liệu của Bob/Bình mà không xét ai đang gửi yêu cầu. Nhóm đã sửa lỗi bằng cách kiểm tra: "
         "'$isOwner = ($record['owner_username'] === $currentUser['username']);' cùng quyền cán bộ y tế ('$isStaff'). Nếu không khớp, "
         "hệ thống lập tức trả về mã HTTP 403 Forbidden."),

        ("Câu 2: Làm sao công cụ phân biệt được giữa 'Từ chối đúng (DENY)' và 'Lỗi mạng / Cổng đóng'?",
         "Trả lời: Đây là nguyên tắc cốt lõi của đề tài (Phép thử 3). HttpClient trong 'src/HttpClient.php' phân loại kết quả thành 3 nhóm rõ rệt: "
         "PASS, FAIL và ERROR. Nếu cổng đóng (CONNECT_FAILED) hoặc quá thời gian chờ (TIMEOUT), kết quả BẮT BUỘC ghi nhận là ERROR kỹ thuật, "
         "tuyệt đối không bao giờ được coi là máy chủ chặn đúng (DENY). Nhờ đó không xảy ra tình trạng báo cáo đạt giả tạo (False Positive)."),

        ("Câu 3: Phân quyền của công cụ AccessTester kết hợp giữa RBAC và ABAC ra sao?",
         "Trả lời: RBAC (Role-Based) kiểm soát dựa trên vai trò như Admin, Doctor, Member, Auditor. ABAC (Attribute-Based) bổ sung thuộc tính phạm vi: "
         "người dùng có vai trò Member chỉ được thao tác trên tài nguyên có thuộc tính phạm vi 'own' (thuộc sở hữu của chính mình). "
         "Bất kỳ quyền nào chưa được khai báo tường minh trong ma trận quyền permissions mặc định đều bị từ chối (DENY)."),

        ("Câu 4: Công cụ bảo vệ an toàn chống lại các nguy cơ tấn công như thế nào?",
         "Trả lời: Công cụ tích hợp 2 lớp phòng vệ chủ động:\n"
         "1. TargetAllowlist: Chỉ cho phép gửi request tới danh sách IP loopback (127.0.0.1, localhost) và dải cổng được duyệt (8205, 8295), "
         "ngăn chặn hoàn toàn kẻ xấu lợi dụng công cụ để quét cổng nội bộ hoặc tấn công SSRF ra mạng ngoài.\n"
         "2. Redactor: Tự động phát hiện và che giấu toàn bộ chuỗi Bearer token, password CSDL trong nhật ký và báo cáo HTML, không để lọt bí mật ra ngoài.")
    ]

    for q, a in qa_list:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(f"❓ {q}")
        r_q.bold = True
        r_q.font.color.rgb = RGBColor(185, 28, 28)
        
        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_after = Pt(6)
        r_a = p_a.add_run(f"💬 {a}")
        r_a.font.size = Pt(9.5)
        r_a.font.color.rgb = RGBColor(15, 23, 42)

    # -------------------------------------------------------------------------
    # PHẦN V: BẢNG TRA CỨU TÀI KHOẢN & LỆNH ĐIỀU KHIỂN NHANH
    # -------------------------------------------------------------------------
    h5 = doc.add_heading("V. BẢNG TRA CỨU TÀI KHOẢN VÀ LỆNH ĐIỀU KHIỂN NHANH", level=1)
    h5.runs[0].font.color.rgb = RGBColor(3, 105, 161)

    tbl_ref = doc.add_table(rows=6, cols=4)
    tbl_ref.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    ref_headers = ["Tên đăng nhập", "Vai trò (Role)", "Bearer Token kiểm thử", "Họ tên & Quyền hạn"]
    for i, h in enumerate(ref_headers):
        cell = tbl_ref.cell(0, i)
        set_cell_background(cell, "1E293B")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    ref_data = [
        ("bs_minh", "doctor (Bác sĩ)", "test_token_bs_minh", "BS. CKI Lê Hoàng Minh (Lập bệnh án, xem hồ sơ)"),
        ("bn_an", "member (Bệnh nhân)", "test_token_bn_an", "Bệnh nhân Nguyễn Văn An (Sở hữu bệnh án 1, 2)"),
        ("bn_binh", "member (Bệnh nhân)", "test_token_bn_binh", "Bệnh nhân Trần Thị Bình (Sở hữu bệnh án 3, 4)"),
        ("admin_bv", "admin (Quản trị)", "test_token_admin_bv", "Quản trị viên Bệnh viện (Toàn quyền, quản lý user)"),
        ("tester_sec", "auditor (Kiểm toán)", "test_token_tester_sec", "Kỹ sư Bảo mật Nguyễn Kiểm Thử (Kiểm toán chỉ đọc)")
    ]

    for row_idx, r_items in enumerate(ref_data, start=1):
        for col_idx, item in enumerate(r_items):
            cell = tbl_ref.cell(row_idx, col_idx)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(item)
            r.font.size = Pt(9)
            if col_idx in [0, 2]:
                r.font.name = "Consolas"

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- CHÚC BUỔI BÁO CÁO CỦA NHÓM 13 ĐẠT KẾT QUẢ XUẤT SẮC ---")
    r_end.bold = True
    r_end.font.size = Pt(11)
    r_end.font.color.rgb = RGBColor(2, 132, 199)

    out_docx = r"C:\University\Nam_3_K1\PTPMMNM\btvn\lab5\Huong_Dan_Demo_Bao_Cao.docx"
    try:
        doc.save(out_docx)
        print(f"SUCCESS: Generated {out_docx}")
    except PermissionError:
        alt_docx = r"C:\University\Nam_3_K1\PTPMMNM\btvn\lab5\Huong_Dan_Demo_Bao_Cao_v2.docx"
        doc.save(alt_docx)
        print(f"NOTE: {out_docx} is locked (e.g. by WPS/Word). Successfully saved to {alt_docx}!")

if __name__ == "__main__":
    build_demo_document()
